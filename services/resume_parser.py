"""Resume Parser Service.

Handles text extraction from PDF, DOCX, and raw text input,
and performs heuristic section and contact information detection.
"""

import io
import re
from typing import Dict, Any, List, Optional
import fitz  # PyMuPDF
import docx
from utils.helpers import sanitize_text


class ResumeParsingError(Exception):
    """Custom exception raised when resume parsing fails."""
    pass


class ResumeParser:
    """Parser to extract raw text and structured sections from PDF and DOCX resumes."""

    SECTION_PATTERNS = {
        'summary': [r'summary', r'professional summary', r'profile', r'objective', r'about me'],
        'skills': [r'skills', r'technical skills', r'core competencies', r'technologies', r'skills\s*&\s*abilities'],
        'experience': [r'experience', r'work experience', r'employment history', r'professional experience', r'work history'],
        'projects': [r'projects', r'academic projects', r'personal projects', r'key projects'],
        'education': [r'education', r'academic background', r'qualifications', r'educational details'],
        'certifications': [r'certifications', r'certificates', r'licenses', r'accreditations'],
        'achievements': [r'achievements', r'awards', r'honors', r'accomplishments']
    }

    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> str:
        """Extract plain text from a PDF byte stream using PyMuPDF."""
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            if len(doc) == 0:
                raise ResumeParsingError("The uploaded PDF document has no readable pages.")

            extracted_pages: List[str] = []
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text")
                if text:
                    extracted_pages.append(text)

            full_text = "\n".join(extracted_pages)
            cleaned = sanitize_text(full_text)
            if not cleaned or len(cleaned) < 30:
                raise ResumeParsingError("Could not extract readable text from PDF. It may be scanned or image-only.")
            return cleaned
        except Exception as e:
            if isinstance(e, ResumeParsingError):
                raise e
            raise ResumeParsingError(f"Failed to parse PDF document: {str(e)}")

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        """Extract plain text from a DOCX byte stream using python-docx."""
        try:
            file_stream = io.BytesIO(file_bytes)
            doc = docx.Document(file_stream)
            paragraphs: List[str] = []

            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append(p.text.strip())

            # Also check table content
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n".join(paragraphs)
            cleaned = sanitize_text(full_text)
            if not cleaned or len(cleaned) < 30:
                raise ResumeParsingError("Could not extract readable text from DOCX document.")
            return cleaned
        except Exception as e:
            if isinstance(e, ResumeParsingError):
                raise e
            raise ResumeParsingError(f"Failed to parse DOCX document: {str(e)}")

    @classmethod
    def parse_file(cls, filename: str, file_bytes: bytes) -> Dict[str, Any]:
        """Entry point for parsing uploaded resume files."""
        if not filename or not file_bytes:
            raise ResumeParsingError("No file or empty file content provided.")

        ext = filename.rsplit('.', 1)[-1].lower()
        if ext == 'pdf':
            raw_text = cls.extract_text_from_pdf(file_bytes)
        elif ext == 'docx':
            raw_text = cls.extract_text_from_docx(file_bytes)
        else:
            raise ResumeParsingError(f"Unsupported file format: .{ext}. Please upload a PDF or DOCX file.")

        return cls.parse_raw_text(raw_text)

    @classmethod
    def parse_raw_text(cls, text: str) -> Dict[str, Any]:
        """Parses raw text into structured resume sections and contact metadata."""
        cleaned_text = sanitize_text(text)
        if not cleaned_text or len(cleaned_text) < 20:
            raise ResumeParsingError("Resume text is too short or empty. Please provide detailed resume content.")

        contact_info = cls._extract_contact_info(cleaned_text)
        sections = cls._extract_sections(cleaned_text)

        return {
            "raw_text": cleaned_text,
            "contact": contact_info,
            "sections": sections,
            "word_count": len(cleaned_text.split()),
            "character_count": len(cleaned_text)
        }

    @staticmethod
    def _extract_contact_info(text: str) -> Dict[str, Optional[str]]:
        """Extracts email, phone, candidate name, and web profiles using regex heuristics."""
        # Email pattern
        email_match = re.search(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
        email = email_match.group(0) if email_match else None

        # Phone pattern (supports multiple formats)
        phone_match = re.search(r'(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}\b', text)
        phone = phone_match.group(0) if phone_match else None

        # LinkedIn & GitHub
        linkedin_match = re.search(r'(?:https?://)?(?:www\.)?linkedin\.com/in/([A-Za-z0-9_-]+)', text, re.IGNORECASE)
        linkedin = linkedin_match.group(0) if linkedin_match else None

        github_match = re.search(r'(?:https?://)?(?:www\.)?github\.com/([A-Za-z0-9_-]+)', text, re.IGNORECASE)
        github = github_match.group(0) if github_match else None

        # Candidate Name: typically the first non-empty line
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        name = None
        if lines:
            first_line = lines[0]
            # Check if first line looks like a name (not an email/URL/header)
            if len(first_line) < 50 and not re.search(r'[@:/]', first_line) and not re.search(r'resume|curriculum|cv', first_line, re.IGNORECASE):
                name = first_line

        return {
            "name": name or "Candidate",
            "email": email,
            "phone": phone,
            "linkedin": linkedin,
            "github": github
        }

    @classmethod
    def _extract_sections(cls, text: str) -> Dict[str, str]:
        """Segments resume text into common sections using header markers."""
        lines = text.split('\n')
        sections: Dict[str, List[str]] = {
            'summary': [],
            'skills': [],
            'experience': [],
            'projects': [],
            'education': [],
            'certifications': [],
            'achievements': [],
            'other': []
        }

        current_section = 'other'
        section_regexes = {}
        for sec, patterns in cls.SECTION_PATTERNS.items():
            pattern_str = r'^\s*(?:' + '|'.join(patterns) + r')\s*[:\-_]?\s*$'
            section_regexes[sec] = re.compile(pattern_str, re.IGNORECASE)

        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue

            detected_section = None
            # Check if line matches a known section header
            if len(stripped) < 45:
                for sec, regex in section_regexes.items():
                    if regex.match(stripped):
                        detected_section = sec
                        break

            if detected_section:
                current_section = detected_section
            else:
                sections[current_section].append(stripped)

        # Convert list of lines back to string for each section
        return {sec: "\n".join(lines_list) for sec, lines_list in sections.items() if lines_list}
