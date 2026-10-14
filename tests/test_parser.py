"""Unit tests for Resume Parser and Section Extractor."""

import io
import pytest
import docx
import fitz
from services.resume_parser import ResumeParser, ResumeParsingError
from utils.helpers import get_sample_candidate


def test_parse_raw_text_success():
    sample = get_sample_candidate()
    result = ResumeParser.parse_raw_text(sample["raw_text"])

    assert result is not None
    assert result["word_count"] > 50
    assert result["contact"]["name"] == "Alex Chen"
    assert result["contact"]["email"] == "alex.chen@email.com"
    assert "skills" in result["sections"]
    assert "education" in result["sections"]


def test_parse_raw_text_too_short():
    with pytest.raises(ResumeParsingError):
        ResumeParser.parse_raw_text("Short text")


def test_parse_docx_in_memory():
    # Build a real docx document in memory
    doc = docx.Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Email: jane.doe@example.com | Phone: 555-123-4567")
    doc.add_paragraph("EDUCATION\nBachelor of Science in Data Science, 2025")
    doc.add_paragraph("SKILLS\nPython, SQL, Machine Learning, Git")

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    result = ResumeParser.parse_file("jane_resume.docx", docx_bytes)
    assert result["contact"]["name"] == "Jane Doe"
    assert result["contact"]["email"] == "jane.doe@example.com"
    assert "Python" in result["raw_text"]


def test_parse_pdf_in_memory():
    # Build a real PDF in memory using fitz (PyMuPDF)
    pdf_doc = fitz.open()
    page = pdf_doc.new_page()
    page.insert_text((50, 72), "John Smith\nEmail: john.smith@test.edu\nPhone: (555) 987-6543\nSKILLS: Python, SQL, C++, Docker\nPROJECTS: Built an AI Search Engine")

    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    result = ResumeParser.parse_file("john_resume.pdf", pdf_bytes)
    assert result["contact"]["name"] == "John Smith"
    assert result["contact"]["email"] == "john.smith@test.edu"
    assert "Python" in result["raw_text"]


def test_parse_unsupported_file_extension():
    with pytest.raises(ResumeParsingError) as exc_info:
        ResumeParser.parse_file("resume.txt", b"some text content")
    assert "Unsupported file format" in str(exc_info.value)
